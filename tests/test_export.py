"""Behavioral regression tests for source preservation and independent checking."""
from datetime import datetime
from decimal import Decimal
import hashlib
import importlib.util
import json
from pathlib import Path
import shutil
import tempfile
import unittest
from xml.etree import ElementTree as ET
from zipfile import ZipFile, ZIP_DEFLATED

import openpyxl
from openpyxl.comments import Comment

ROOT = Path(__file__).resolve().parents[1]


def load_script(name):
    spec = importlib.util.spec_from_file_location(name, ROOT / "scripts" / (name + ".py"))
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


exporter = load_script("export_workbook")
verifier = load_script("verify_export")
NS = {"s": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}


def make_workbook(path):
    workbook = openpyxl.Workbook()
    shipments = workbook.active
    shipments.title = "Бензовозы"
    stock = workbook.create_sheet("Склад")
    payments = workbook.create_sheet("Выписка")
    for ref, value in {"C13": "Дата", "D13": "Контрагент", "M13": "Поставщик"}.items():
        shipments[ref] = value
    for ref, value in {"B1": "Дата", "C1": "Поступление", "D1": "Списание", "F1": "Контрагент"}.items():
        payments[ref] = value
    shipments["A1"].comment = Comment("Комментарий пустой ячейки", "Проверка")
    shipments["C15"] = datetime(2026, 9, 1)
    shipments["D15"] = 'ООО "Альфа"'
    shipments["M15"] = "Поставщик"
    shipments["E15"] = "Иван"
    shipments["L15"] = 1
    shipments["J15"] = "=L15/#REF!"
    shipments["W15"] = "Сохранить примечание"
    shipments["C16"] = datetime(2026, 9, 2)
    shipments["D16"] = '  ооо   "АЛЬФА"  '
    shipments["M16"] = "Поставщик"
    shipments["E16"] = "иван"
    shipments["L16"] = 0.02
    # Formula-only rows must remain in raw; they are not real shipments/payments.
    shipments["B17"] = '=IF(C17="","",MONTH(C17))'
    shipments["L17"] = "=1+1"
    payments["B2"] = datetime(2026, 9, 1)
    payments["C2"] = "1\u00a0234,56"
    payments["F2"] = 'ООО "Альфа"'
    payments["B3"] = datetime(2026, 9, 2)  # Incomplete date-only source row is retained.
    payments["A4"] = '=IF(B4="","",MONTH(B4))'
    payments["C4"] = "=1+1"
    payments["F5"] = "Ошибка в формуле"
    payments["C5"] = "=1/0"
    payments["B6"] = datetime(2026, 9, 3)
    payments["D6"] = 12.01
    payments["F6"] = "Поставщик"
    stock["B1"] = "Склад"
    stock["B4"] = "=1/0"
    workbook.save(path)
    workbook.close()
    # Inject exact serialized numeric and cached error evidence directly into OOXML.
    with ZipFile(path) as original:
        files = {name: original.read(name) for name in original.namelist()}
    for name, replacements in {
        "xl/worksheets/sheet1.xml": {"L15": ("n", "123456789.123456789001")},
        "xl/worksheets/sheet2.xml": {"B4": ("e", "#DIV/0!")},
        "xl/worksheets/sheet3.xml": {"C5": ("e", "#DIV/0!")},
    }.items():
        tree = ET.fromstring(files[name])
        for cell in tree.findall(".//s:c", NS):
            if cell.attrib["r"] in replacements:
                kind, value = replacements[cell.attrib["r"]]
                cell.set("t", kind)
                v = cell.find("s:v", NS)
                if v is None:
                    v = ET.SubElement(cell, "{" + NS["s"] + "}v")
                v.text = value
        files[name] = ET.tostring(tree, encoding="utf-8", xml_declaration=True)
    with ZipFile(path, "w", ZIP_DEFLATED) as destination:
        for name, value in files.items():
            destination.writestr(name, value)


class MoneyTests(unittest.TestCase):
    def test_exact_decimal_values_and_russian_formats(self):
        examples = {"123456789.123456789001": Decimal("123456789.123456789001"),
                    "1e3": Decimal("1000"), "-1.2E-9": Decimal("-0.0000000012"),
                    "1 234,56": Decimal("1234.56"), "1\u00a0234,56": Decimal("1234.56"),
                    "1\u202f234,56": Decimal("1234.56"), "-0,01": Decimal("-0.01"),
                    "+12,5": Decimal("12.5"), "0": Decimal(0)}
        for source, expected in examples.items():
            with self.subTest(source=source):
                self.assertEqual(exporter.parse_money(source), expected)
        self.assertEqual(exporter.parse_money(Decimal("1.000000001")), Decimal("1.000000001"))

    def test_malformed_values_are_not_silently_repaired(self):
        for source in ("1_000", "12 34,56", "1,234", "1,23,45", "1.234,56", "1  234,56", "12 руб.",
                       "NaN", "Infinity", "-Infinity", "", " ", None, True, False, [], {}):
            with self.subTest(source=source):
                self.assertIsNone(exporter.parse_money(source))


class ExportTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.temp = tempfile.TemporaryDirectory(prefix="artel-export-tests-")
        cls.directory = Path(cls.temp.name)
        cls.source = cls.directory / "fixture.xlsx"
        cls.output = cls.directory / "snapshot"
        make_workbook(cls.source)
        cls.original_hash = hashlib.sha256(cls.source.read_bytes()).hexdigest()
        cls.report = exporter.export_workbook(cls.source, cls.output)

    @classmethod
    def tearDownClass(cls):
        cls.temp.cleanup()

    def data(self, name):
        return json.loads((self.output / (name + ".json")).read_text(encoding="utf-8"))["data"]

    def test_export_does_not_modify_source_and_refuses_overwrite(self):
        self.assertEqual(hashlib.sha256(self.source.read_bytes()).hexdigest(), self.original_hash)
        before = (self.output / "manifest.json").read_bytes()
        with self.assertRaises(ValueError):
            exporter.export_workbook(self.source, self.output)
        self.assertEqual((self.output / "manifest.json").read_bytes(), before)

    def test_transaction_rows_exclude_templates_but_keep_incomplete_evidence(self):
        self.assertEqual([r["source"]["row"] for r in self.data("shipments")], [15, 16])
        self.assertEqual([r["source"]["row"] for r in self.data("payments")], [2, 3, 5, 6])
        self.assertIn(17, self.report["excluded_from_transaction_lists"]["shipment_template_rows"])
        self.assertIn(4, self.report["excluded_from_transaction_lists"]["payment_template_rows"])
        rows = {r["source"]["row"]: r for r in self.data("payments")}
        self.assertEqual(rows[3]["record_kind"], "incomplete_source_row")
        self.assertIn("no_nonzero_amount", rows[3]["quality_flags"])
        self.assertEqual(rows[2]["normalized_amounts"]["incoming_amount"], "1234.56")
        self.assertEqual(rows[5]["normalized_amounts"]["incoming_amount"], None)

    def test_preserves_formula_error_precision_empty_comments_and_notes(self):
        raw = {s["name"]: s["cells"] for s in self.data("raw_workbook")["sheets"]}
        self.assertEqual(raw["Бензовозы"]["L15"]["value"], "123456789.123456789001")
        self.assertEqual(raw["Бензовозы"]["A1"]["comment"]["text"], "Комментарий пустой ячейки")
        self.assertEqual(raw["Бензовозы"]["A1"]["value"], None)
        self.assertEqual(raw["Бензовозы"]["J15"]["formula"], "=L15/#REF!")
        self.assertEqual(raw["Бензовозы"]["L17"]["value"], None)
        self.assertEqual(raw["Бензовозы"]["L17"]["cache_status"], "empty_or_missing")
        self.assertEqual(raw["Выписка"]["C5"]["value"], "#DIV/0!")
        self.assertEqual(raw["Выписка"]["C5"]["value_type"], "error")
        self.assertEqual(raw["Выписка"]["C5"]["formula"], "=1/0")
        self.assertEqual(self.data("shipments")[0]["fields"]["unlabelled_note"], "Сохранить примечание")
        self.assertEqual(Decimal(self.report["source_totals"]["shipments"]["customer_amount"]["total"]), Decimal("123456789.143456789001"))
        self.assertEqual(Decimal(self.report["source_totals"]["payments"]["incoming_amount"]["total"]), Decimal(0))
        self.assertEqual(Decimal(self.report["normalized_payment_totals"]["incoming_amount"]), Decimal("1234.56"))

    def test_labels_never_create_accounts_or_assign_owners(self):
        companies = self.data("companies")
        alpha = next(c for c in companies if "альфа" in c["normalized_name"])
        self.assertEqual(len(alpha["source_names"]), 2)
        self.assertNotIn("multiple_source_manager_labels", alpha["quality_flags"])
        for item in companies + self.data("shipments") + self.data("payments"):
            self.assertIsNone(item["owner_user_id"])
        managers = self.data("manager_labels")
        self.assertEqual(len(managers), 1)
        self.assertFalse(managers[0]["is_user_account"])
        self.assertIsNone(managers[0]["user_id"])
        self.assertEqual(len(managers[0]["shipment_ids"]), 2)
        roles = json.loads((ROOT / "config/roles.json").read_text(encoding="utf-8"))
        self.assertEqual(roles["status"], "specification_only")
        self.assertFalse(roles["authorization_implemented"])
        self.assertEqual(roles["accounts"], [])
        self.assertEqual(roles["ownership_assignments"], [])

    def test_independent_verifier_accepts_valid_fixture(self):
        result = verifier.verify(self.source, self.output, ROOT / "config/roles.json")
        self.assertEqual(result["status"], "passed", result["failures"])

    def changed_snapshot(self, name, edit, repair_manifest=True):
        temp = tempfile.TemporaryDirectory(prefix="artel-corrupt-test-")
        self.addCleanup(temp.cleanup)
        target = Path(temp.name) / "snapshot"
        shutil.copytree(self.output, target)
        file = target / (name + ".json")
        payload = json.loads(file.read_text(encoding="utf-8"))
        edit(payload["data"])
        file.write_text(json.dumps(payload, ensure_ascii=False), encoding="utf-8")
        if repair_manifest:
            manifest_file = target / "manifest.json"
            manifest = json.loads(manifest_file.read_text(encoding="utf-8"))
            manifest["files"][file.name] = {"sha256": hashlib.sha256(file.read_bytes()).hexdigest(), "bytes": file.stat().st_size}
            manifest_file.write_text(json.dumps(manifest, ensure_ascii=False), encoding="utf-8")
        return verifier.verify(self.source, target, ROOT / "config/roles.json")

    def test_verifier_detects_rehashed_source_value_corruption(self):
        def corrupt(raw):
            next(s for s in raw["sheets"] if s["name"] == "Бензовозы")["cells"]["L15"]["value"] = "123456789.12"
        result = self.changed_snapshot("raw_workbook", corrupt)
        self.assertEqual(result["status"], "failed")
        self.assertGreater(result["failures_by_category"].get("source_cells", 0), 0)

    def test_verifier_detects_rehashed_owner_assignment(self):
        result = self.changed_snapshot("shipments", lambda rows: rows[0].update(owner_user_id="invented-user"))
        self.assertGreater(result["failures_by_category"].get("ownership", 0), 0)

    def test_verifier_detects_rehashed_total_corruption(self):
        def corrupt(report):
            report["normalized_payment_totals"]["incoming_amount"] = "1234.55"
        result = self.changed_snapshot("validation_report", corrupt)
        self.assertGreater(result["failures_by_category"].get("decimal_totals", 0), 0)

    def test_verifier_detects_manifest_tampering(self):
        result = self.changed_snapshot("payments", lambda rows: rows[0]["quality_flags"].append("changed"), repair_manifest=False)
        self.assertGreater(result["failures_by_category"].get("hashes", 0), 0)

    def test_unknown_source_headers_fail_before_export(self):
        source = self.directory / "wrong-layout.xlsx"
        make_workbook(source)
        workbook = openpyxl.load_workbook(source)
        workbook["Бензовозы"]["D13"] = "Другой столбец"
        workbook.save(source)
        workbook.close()
        target = self.directory / "wrong-layout-output"
        with self.assertRaises(ValueError):
            exporter.export_workbook(source, target)
        self.assertFalse(target.exists())


if __name__ == "__main__":
    unittest.main()
