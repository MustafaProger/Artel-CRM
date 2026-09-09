#!/usr/bin/env python3
"""Local, source-preserving Artel workbook export. Never evaluates formulas."""
from __future__ import annotations

import argparse
from collections import Counter, defaultdict
from datetime import date, datetime, timezone
from decimal import Decimal, InvalidOperation
from difflib import SequenceMatcher
import hashlib
import json
from pathlib import Path
import posixpath
import re
import shutil
import tempfile
import unicodedata
from uuid import NAMESPACE_URL, uuid5
from xml.etree import ElementTree as ET
from zipfile import ZipFile

import openpyxl

NS = {"s": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
SHIPMENT_FIELDS = {
    "A": "document_number", "B": "month", "C": "date", "D": "customer_name",
    "E": "manager_label", "F": "payment_form", "G": "product",
    "H": "quantity_tonnes", "I": "quantity_litres", "J": "sale_price_per_tonne",
    "K": "sale_price_per_litre", "L": "customer_amount", "M": "supplier_name",
    "N": "purchase_price_unspecified_unit", "O": "purchase_amount",
    "P": "carrier_name", "Q": "transport_amount", "R": "kvp_source",
    "S": "profit_source", "T": "paid_amount_source", "U": "debt_overpayment_source",
    "V": "term_source", "W": "unlabelled_note",
}
PAYMENT_FIELDS = {"A": "month", "B": "date", "C": "incoming_amount",
                  "D": "outgoing_amount", "E": "purpose", "F": "counterparty_name",
                  "G": "unlabelled_extra"}


def normal_name(value):
    return " ".join(unicodedata.normalize("NFKC", str(value)).split()).casefold()


def stable_id(kind, key):
    return str(uuid5(NAMESPACE_URL, f"artel-crm:{kind}:{key}"))


def number(value):
    try:
        if value is None or isinstance(value, bool):
            return None
        text = str(value).strip()
        if not re.fullmatch(r"[+-]?(?:\d+(?:\.\d*)?|\.\d+)(?:[Ee][+-]?\d+)?", text):
            return None
        result = Decimal(text)
        return result if result.is_finite() else None
    except InvalidOperation:
        return None


def decimal_text(value):
    return format(value, "f")


def parse_money(value):
    """Only explicitly supported decimal and Russian formatted amount syntax."""
    direct = number(value)
    if direct is not None:
        return direct
    if not isinstance(value, str):
        return None
    stripped = value.strip().replace("\u00a0", " ").replace("\u202f", " ")
    if re.fullmatch(r"[+-]?(?:\d{1,3}(?: \d{3})+|\d+)(?:,\d{1,2})?", stripped):
        return Decimal(stripped.replace(" ", "").replace(",", "."))
    return None


def write_json(path, value):
    path.write_text(json.dumps(value, ensure_ascii=False, indent=2, allow_nan=False) + "\n", encoding="utf-8")


def xml_sheets(path):
    """Read exact numeric lexemes instead of round-tripping through binary floats."""
    result = {}
    with ZipFile(path) as archive:
        rels = ET.fromstring(archive.read("xl/_rels/workbook.xml.rels"))
        targets = {r.attrib["Id"]: r.attrib["Target"] for r in rels}
        workbook = ET.fromstring(archive.read("xl/workbook.xml"))
        for sheet in workbook.find("s:sheets", NS):
            target = targets[sheet.attrib["{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id"]]
            xml_path = target.lstrip("/") if target.startswith("/") else posixpath.normpath(posixpath.join("xl", target))
            root = ET.fromstring(archive.read(xml_path))
            cells = {}
            for c in root.findall(".//s:sheetData/s:row/s:c", NS):
                v, f, inline = c.find("s:v", NS), c.find("s:f", NS), c.find("s:is", NS)
                if v is None and f is None and inline is None:
                    continue  # A style-only cell does not carry data.
                cells[c.attrib["r"]] = {
                    "xml_type": c.attrib.get("t", "n"),
                    "xml_value": v.text if v is not None else None,
                    "xml_value_present": v is not None,
                    "formula_xml": f.text if f is not None else None,
                    "formula_attributes": dict(f.attrib) if f is not None else None,
                    "has_formula": f is not None,
                }
            result[sheet.attrib["name"]] = cells
    return result


def extract_raw(path):
    exact = xml_sheets(path)
    formulas = openpyxl.load_workbook(path, data_only=False)
    cached = openpyxl.load_workbook(path, data_only=True)
    sheets, lookup, issues = [], {}, []
    for ws in formulas:
        for row in ws:
            for cell in row:
                if cell.comment or cell.hyperlink:
                    exact[ws.title].setdefault(cell.coordinate, {
                        "xml_type": "n", "xml_value": None, "xml_value_present": False,
                        "formula_xml": None, "formula_attributes": None, "has_formula": False,
                    })
        entries = {}
        for address, xml in exact[ws.title].items():
            cell, cache = ws[address], cached[ws.title][address]
            value = cache.value if xml["has_formula"] else cell.value
            kind = "empty"
            if isinstance(value, (datetime, date)):
                kind, value = "datetime", value.isoformat()
            elif isinstance(value, bool):
                kind = "boolean"
            elif (cache if xml["has_formula"] else cell).data_type == "e":
                kind = "error"
            elif xml["xml_type"] == "n" and xml["xml_value"] is not None:
                kind, value = "decimal", xml["xml_value"]
            elif value is not None:
                kind, value = "string", str(value)
            elif xml["xml_type"] == "str" and xml["xml_value_present"]:
                kind, value = "string", ""
            formula = str(cell.value) if xml["has_formula"] else None
            entry = {"value": value, "value_type": kind, "formula": formula,
                     "number_format": cell.number_format, **xml}
            if xml["has_formula"]:
                entry["cache_status"] = (
                    "error" if kind == "error" else
                    "present" if value is not None else
                    "empty_or_missing")
            if cell.comment:
                entry["comment"] = {"text": cell.comment.text, "author": cell.comment.author}
            if cell.hyperlink:
                entry["hyperlink"] = {"target": cell.hyperlink.target, "location": cell.hyperlink.location}
            entries[address] = entry
            if formula and "#REF!" in formula.upper():
                issues.append({"code": "broken_formula_reference", "severity": "error",
                               "sheet": ws.title, "cell": address})
            if kind == "error":
                issues.append({"code": "cached_excel_error", "severity": "error",
                               "sheet": ws.title, "cell": address, "value": value})
            if formula and value is None:
                issues.append({"code": "formula_cache_empty_or_missing", "severity": "warning",
                               "sheet": ws.title, "cell": address})
        sheets.append({"name": ws.title, "state": ws.sheet_state,
                       "max_row": ws.max_row, "max_column": ws.max_column,
                       "merged_ranges": [str(r) for r in ws.merged_cells.ranges],
                       "hidden_rows": [i for i, d in ws.row_dimensions.items() if d.hidden],
                       "hidden_columns": [i for i, d in ws.column_dimensions.items() if d.hidden],
                       "auto_filter": ws.auto_filter.ref, "cells": entries})
        lookup[ws.title] = entries
    calc = formulas.calculation
    return {"sheets": sheets, "calculation_properties": dict(calc) if calc else {}}, lookup, issues


def export_workbook(source: Path, output: Path, source_url=None):
    source = source.resolve()
    output = output.resolve()
    if output.exists():
        raise ValueError(f"Output already exists: {output}. Use a new snapshot directory.")
    source_hash = hashlib.sha256(source.read_bytes()).hexdigest()
    raw, lookup, issues = extract_raw(source)
    required = {"Бензовозы", "Склад", "Выписка"}
    if not required.issubset(lookup):
        raise ValueError("Expected sheets: Бензовозы, Склад, Выписка. Source layout requires review.")
    for sheet, refs in {"Бензовозы": {"C13": "Дата", "D13": "Контрагент", "M13": "Поставщик"},
                        "Выписка": {"B1": "Дата", "C1": "Поступление", "D1": "Списание", "F1": "Контрагент"}}.items():
        for ref, expected in refs.items():
            if normal_name(lookup[sheet].get(ref, {}).get("value", "")) != normal_name(expected):
                raise ValueError(f"Unexpected header at {sheet}!{ref}; review mapping before importing.")

    companies = {}
    managers = {}

    def value(sheet, address):
        return lookup[sheet].get(address, {}).get("value")

    def register(name, role, sheet, address, record_id=None):
        if name is None or not str(name).strip():
            return None
        key = normal_name(name)
        cid = stable_id("counterparty", key)
        item = companies.setdefault(cid, {
            "id": cid, "display_name": " ".join(str(name).split()),
            "normalized_name": key, "source_names": [], "source_roles": [],
            "entity_kind": "internal_location" if key == "склад" else "unverified_counterparty",
            "inn": None, "kpp": None, "ogrn": None, "legal_address": None,
            "contacts": [], "owner_user_id": None,
            "verification_status": "source_only_not_registry_verified",
            "occurrences": [], "shipment_ids": [], "payment_ids": [], "manager_labels": [],
        })
        if str(name) not in item["source_names"]:
            item["source_names"].append(str(name))
        if role not in item["source_roles"]:
            item["source_roles"].append(role)
        item["occurrences"].append({"sheet": sheet, "cell": address, "role": role, "record_id": record_id})
        if record_id:
            field = "payment_ids" if role == "payment_counterparty" else "shipment_ids"
            if record_id not in item[field]:
                item[field].append(record_id)
        return cid

    def source_rows(sheet, start, literal_columns):
        rows = defaultdict(dict)
        for address, cell in lookup[sheet].items():
            col, row = re.fullmatch(r"([A-Z]+)(\d+)", address).groups()
            if int(row) >= start:
                rows[int(row)][col] = cell
        active, templates = [], []
        for row, cells in sorted(rows.items()):
            literal = any(col in literal_columns and not c["has_formula"] and c["value"] is not None
                          and str(c["value"]).strip() for col, c in cells.items())
            if literal:
                active.append((row, cells))
            else:
                templates.append(row)
        return active, templates

    shipment_rows, shipment_templates = source_rows("Бензовозы", 15, set(SHIPMENT_FIELDS) - {"B"})
    payment_rows, payment_templates = source_rows("Выписка", 2, set(PAYMENT_FIELDS) - {"A"})
    for sheet, cells in lookup.items():
        for address, cell in cells.items():
            formula = cell.get("formula") or ""
            for target_sheet, last_data_row in [("Бензовозы", max((r for r, _ in shipment_rows), default=14)),
                                                ("Выписка", max((r for r, _ in payment_rows), default=1))]:
                # A warning about coverage, not a claim that every restricted range is wrong.
                if sheet != "Склад" and not (sheet == "Бензовозы" and int(re.search(r"\d+", address)[0]) < 12):
                    continue
                pattern = rf"'?{re.escape(target_sheet)}'?!\$?[A-Z]+\$?\d+:\$?[A-Z]+\$?(\d+)"
                ends = sorted({int(end) for end in re.findall(pattern, formula) if int(end) < last_data_row})
                if ends:
                    issues.append({"code": "summary_range_ends_before_last_source_row", "severity": "warning",
                                   "sheet": sheet, "cell": address, "target_sheet": target_sheet,
                                   "range_end_rows": ends, "last_source_row": last_data_row})

    def record(sheet, row, cells, fields, kind):
        rid = stable_id(kind, f"{source_hash}:{sheet}:{row}")
        result = {"id": rid, "source": {"sheet": sheet, "row": row, "workbook_sha256": source_hash},
                  "owner_user_id": None, "fields": {field: cells.get(col, {}).get("value") for col, field in fields.items()},
                  "cells": {col + str(row): c for col, c in cells.items()},
                  "quality_flags": []}
        date_value = result["fields"]["date"]
        if not date_value or not re.match(r"^\d{4}-\d{2}-\d{2}", str(date_value)):
            result["quality_flags"].append("missing_or_invalid_date")
        result["record_status"] = "requires_review" if result["quality_flags"] else "source_record"
        return result

    shipments = []
    for row, cells in shipment_rows:
        item = record("Бензовозы", row, cells, SHIPMENT_FIELDS, "shipment")
        for col, field, role in [("D", "customer_id", "customer"), ("M", "supplier_id", "supplier"), ("P", "carrier_id", "carrier")]:
            item[field] = register(value("Бензовозы", f"{col}{row}"), role, "Бензовозы", f"{col}{row}", item["id"])
        if item["customer_id"] is None:
            item["quality_flags"].append("missing_customer")
        if item["supplier_id"] is None:
            item["quality_flags"].append("missing_supplier")
        m = cells.get("E", {})
        label = item["fields"]["manager_label"]
        if m.get("formula") and "#REF!" in m["formula"]:
            item["quality_flags"].append("manager_formula_broken")
        if m.get("value_type") == "string" and label and str(label).strip():
            mid = stable_id("manager_label", normal_name(label))
            managers.setdefault(mid, {"id": mid, "source_label": str(label), "user_id": None,
                                      "is_user_account": False, "shipment_ids": []})["shipment_ids"].append(item["id"])
            if item["customer_id"] and str(label) not in companies[item["customer_id"]]["manager_labels"]:
                companies[item["customer_id"]]["manager_labels"].append(str(label))
        else:
            item["fields"]["manager_label"] = None
            item["quality_flags"].append("manager_unresolved")
        if item["quality_flags"]:
            item["record_status"] = "requires_review"
        shipments.append(item)

    payments = []
    for row, cells in payment_rows:
        item = record("Выписка", row, cells, PAYMENT_FIELDS, "payment")
        item["counterparty_id"] = register(value("Выписка", f"F{row}"), "payment_counterparty", "Выписка", f"F{row}", item["id"])
        if not item["counterparty_id"]:
            item["quality_flags"].append("missing_counterparty")
        item["normalized_amounts"] = {}
        for col, field in [("C", "incoming_amount"), ("D", "outgoing_amount")]:
            cell = cells.get(col, {})
            parsed = parse_money(item["fields"][field])
            item["normalized_amounts"][field] = decimal_text(parsed) if parsed is not None else None
            if cell.get("value_type") == "string" and parsed is not None:
                item["quality_flags"].append(f"{field}_stored_as_text")
                issues.append({"code": "money_stored_as_text", "severity": "warning", "sheet": "Выписка",
                               "cell": f"{col}{row}", "source_value": cell["value"], "parsed_decimal": decimal_text(parsed)})
            elif cell.get("value") is not None and parsed is None:
                item["quality_flags"].append(f"{field}_not_parseable")
        incoming, outgoing = (parse_money(item["fields"][f]) for f in ("incoming_amount", "outgoing_amount"))
        if incoming and outgoing:
            item["quality_flags"].append("both_incoming_and_outgoing")
        if not incoming and not outgoing:
            item["quality_flags"].append("no_nonzero_amount")
        item["record_kind"] = "payment" if incoming is not None or outgoing is not None else "incomplete_source_row"
        if item["quality_flags"]:
            item["record_status"] = "requires_review"
        payments.append(item)

    summary_refs = [("D5", "prepayment_summary", ["E5", "F5"])]
    summary_refs += [(f"M{r}", "supplier_summary", [f"N{r}"]) for r in range(2, 12)]
    summary_refs += [(f"P{r}", "supplier_summary", [f"R{r}"]) for r in range(2, 12)]
    summary_refs += [(f"T{r}", "supplier_summary", [f"V{r}"]) for r in range(2, 12)]
    summaries = []
    for address, role, related in summary_refs:
        cid = register(value("Бензовозы", address), role, "Бензовозы", address)
        if cid:
            summaries.append({"counterparty_id": cid, "source": {"sheet": "Бензовозы", "cell": address},
                              "role": role, "cells": {a: lookup["Бензовозы"].get(a) for a in [address, *related]}})

    # Unlabelled source cells are separate evidence, never an automatic alias relation.
    for row, cells in payment_rows:
        extra = cells.get("G", {})
        if extra.get("value_type") == "string" and extra.get("value"):
            register(extra["value"], "unlabelled_payment_note", "Выписка", f"G{row}")

    stock = []
    for title_cell, month_col, columns, start in [
        ("B1", "A", "BCDEFG", 4), ("J1", "I", "JKLMNO", 4),
        ("B16", "A", "BCDEFG", 19), ("J16", "I", "JKLMNO", 19),
        ("B31", "A", "BCDEFG", 34),
    ]:
        cid = register(value("Склад", title_cell), "stock_summary", "Склад", title_cell)
        for row in range(start, start + 12):
            field_names = ["incoming_litres", "incoming_amount", "outgoing_litres", "outgoing_amount", "balance_litres", "balance_amount"]
            refs = [f"{c}{row}" for c in columns]
            stock.append({"counterparty_id": cid, "label": value("Склад", title_cell),
                          "month": value("Склад", f"{month_col}{row}"),
                          "source": {"sheet": "Склад", "row": row, "title_cell": title_cell},
                          "fields": dict(zip(field_names, [value("Склад", a) for a in refs])),
                          "cells": {a: lookup["Склад"].get(a) for a in refs},
                          "value_basis": "xlsx_cached_not_recalculated"})

    duplicate_rows = []
    for label, records in [("shipments", shipments), ("payments", payments)]:
        comparison_columns = list("CDFGHIJKLMNO PQRT".replace(" ", "")) if label == "shipments" else list("BCDEF")
        mapping = SHIPMENT_FIELDS if label == "shipments" else PAYMENT_FIELDS
        comparison_fields = [mapping[c] for c in comparison_columns]
        groups = defaultdict(list)
        for item in records:
            # Matching business fields is only a candidate, never deletion authorization.
            key = json.dumps({f: item["fields"][f] for f in comparison_fields}, sort_keys=True, ensure_ascii=False)
            groups[key].append(item["source"]["row"])
        duplicate_rows.extend({"dataset": label, "rows": rows, "comparison_fields": comparison_fields, "status": "candidate_not_removed"}
                              for rows in groups.values() if len(rows) > 1)

    candidates = []
    entries = sorted(companies.values(), key=lambda x: x["normalized_name"])
    for i, first in enumerate(entries):
        for second in entries[i + 1:]:
            a, b = first["normalized_name"], second["normalized_name"]
            similarity = SequenceMatcher(None, a, b).ratio()
            legal_key = lambda name: re.sub(r"[^\w]", "", re.sub(r"\b(ооо|пао|ао|ип|оао|зао)\b", "", name))
            if similarity >= 0.9 or (len(legal_key(a)) >= 4 and legal_key(a) == legal_key(b)):
                candidates.append({"ids": [first["id"], second["id"]],
                                   "names": [first["display_name"], second["display_name"]],
                                   "similarity": round(similarity, 3), "action": "manual_review_no_merge"})
    for item in entries:
        item["inn_candidates"] = [
            {"value": match, "source_name": name, "status": "embedded_in_name_not_registry_verified"}
            for name in item["source_names"] for match in re.findall(r"(?<!\d)(\d{12}|\d{10})(?!\d)", name)
        ]
        item["quality_flags"] = ["legal_identity_unverified", "owner_unassigned"]
        if len(item["source_names"]) > 1:
            item["quality_flags"].append("spacing_or_case_variants_grouped")
        if len({normal_name(label) for label in item["manager_labels"]}) > 1:
            item["quality_flags"].append("multiple_source_manager_labels")

    alias_groups = defaultdict(list)
    for item in entries:
        key = re.sub(r"[^\w]", "", re.sub(r"\b(ооо|пао|ао|ип|оао|зао)\b", "", item["normalized_name"]))
        alias_groups[key].append(item["id"])

    def totals(records, fields):
        result = {}
        for field in fields:
            values = []
            for item in records:
                mapping = SHIPMENT_FIELDS if item["source"]["sheet"] == "Бензовозы" else PAYMENT_FIELDS
                column = next(c for c, f in mapping.items() if f == field)
                cell = item["cells"].get(f"{column}{item['source']['row']}", {})
                values.append(number(cell.get("value")) if cell.get("value_type") == "decimal" else None)
            result[field] = {"total": decimal_text(sum((v for v in values if v is not None), Decimal(0))),
                             "numeric_cells": sum(v is not None for v in values),
                             "empty_or_non_numeric_cells": sum(v is None for v in values)}
        return result

    unmapped_evidence = [
        {"sheet": sheet, "cell": address, "kind": kind, "source_value": cell.get("value"),
         **({"comment": cell["comment"]} if kind == "comment" else {})}
        for sheet, cells in lookup.items() for address, cell in cells.items()
        for kind in (["comment"] if cell.get("comment") else []) +
                    (["unlabelled_column"] if (sheet == "Бензовозы" and address.startswith("W")) or
                     (sheet == "Выписка" and address.startswith("G") and address != "G1") else [])
    ]
    identifier_mentions = [
        {"sheet": "Выписка", "cell": f"E{item['source']['row']}", "payment_id": item["id"],
         "candidate_inn": match, "status": "mentioned_in_purpose_not_assigned_to_counterparty"}
        for item in payments
        for match in re.findall(r"ИНН\s*[:№]?\s*(\d{12}|\d{10})(?!\d)", str(item["fields"]["purpose"]), re.I)
    ]

    report = {
        "status": "exported_with_source_issues", "google_verified": False, "registry_verified": False,
        "counts": {"sheets": len(raw["sheets"]), "raw_cells": sum(len(s["cells"]) for s in raw["sheets"]),
                   "formula_cells": sum(c["has_formula"] for s in raw["sheets"] for c in s["cells"].values()),
                   "counterparties": len(entries), "shipment_rows": len(shipments), "payment_rows": len(payments),
                   "payments_with_amount": sum(i["record_kind"] == "payment" for i in payments),
                   "incomplete_payment_rows": sum(i["record_kind"] == "incomplete_source_row" for i in payments),
                   "manager_labels": len(managers), "stock_monthly_rows": len(stock),
                   "company_summary_rows": len(summaries)},
        "excluded_from_transaction_lists": {"reason": "no nonempty business literals; retained in raw_workbook.json",
                                             "shipment_template_rows": shipment_templates, "payment_template_rows": payment_templates},
        "issue_counts": dict(Counter(i["code"] for i in issues)),
        "cell_issues": issues,
        "record_flag_counts": {label: dict(Counter(f for i in records for f in i["quality_flags"]))
                               for label, records in [("shipments", shipments), ("payments", payments)]},
        "duplicate_record_candidates": duplicate_rows, "similar_counterparty_candidates": candidates,
        "legal_form_alias_candidate_groups": [{"comparison_key": key, "ids": ids, "action": "manual_review_no_merge"}
                                               for key, ids in alias_groups.items() if len(ids) > 1],
        "multiple_manager_companies": [i["id"] for i in entries if "multiple_source_manager_labels" in i["quality_flags"]],
        "unmapped_evidence": unmapped_evidence, "identifier_mentions": identifier_mentions,
        "source_totals": {"basis": "Exact decimal sum of saved XLSX values, not verified balances or recalculated formulas.",
                          "shipments": totals(shipments, ["quantity_tonnes", "quantity_litres", "customer_amount", "purchase_amount", "transport_amount", "paid_amount_source", "profit_source"]),
                          "payments": totals(payments, ["incoming_amount", "outgoing_amount"])},
        "normalized_payment_totals": {
            field: decimal_text(sum((parse_money(i["fields"][field]) or Decimal(0) for i in payments), Decimal(0)))
            for field in ["incoming_amount", "outgoing_amount"]
        },
        "limitations": ["No live Google verification without source link/access.",
                        "All formulas and their saved values preserved; formulas are not executed.",
                        "Missing cache is not zero; source spreadsheet must be recalculated for current values.",
                        "Names are source labels; no legal entity identity or registry status inferred.",
                        "No INN/contact directory or authenticated user mapping is available in mapped columns; identifiers embedded in names are candidates only.",
                        "No manager ownership assigned automatically.",
                        "IDs of source transactions are snapshot-specific; future sync needs an explicit reconciliation policy."],
    }
    meta = {"schema_version": "1.0.0", "created_at_utc": datetime.now(timezone.utc).isoformat(),
            "source_file": source.name, "source_sha256": source_hash,
            "source_kind": "local_xlsx", "google_source_url": source_url,
            "google_verified": False, "decimal_encoding": "string",
            "formula_policy": "preserve formula and saved result, never evaluate",
            "ownership_policy": "unassigned until explicit authenticated-user mapping"}
    output.parent.mkdir(parents=True, exist_ok=True)
    stage = Path(tempfile.mkdtemp(prefix=".export-", dir=output.parent))
    try:
        datasets = {"raw_workbook": raw, "companies": entries, "shipments": shipments,
                    "payments": payments, "manager_labels": list(managers.values()),
                    "stock_summaries": stock, "company_summaries": summaries, "validation_report": report}
        for filename, data in datasets.items():
            write_json(stage / f"{filename}.json", {"meta": meta, "data": data})
        manifest = {"meta": meta, "counts": report["counts"], "files": {}}
        for file in sorted(stage.glob("*.json")):
            manifest["files"][file.name] = {"sha256": hashlib.sha256(file.read_bytes()).hexdigest(), "bytes": file.stat().st_size}
        write_json(stage / "manifest.json", manifest)
        stage.rename(output)
    except BaseException:
        shutil.rmtree(stage)
        raise
    return report


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("source", type=Path)
    parser.add_argument("--output", type=Path, required=True, help="New snapshot directory; existing snapshots are never overwritten")
    parser.add_argument("--source-url", help="Optional provenance only; this flag does not fetch or verify Google Sheets")
    args = parser.parse_args()
    result = export_workbook(args.source, args.output, args.source_url)
    print(json.dumps({"output": str(args.output.resolve()), "counts": result["counts"], "issue_counts": result["issue_counts"]}, ensure_ascii=False, indent=2))
