#!/usr/bin/env python3
"""Independently verify an Artel JSON snapshot against the original XLSX.

This verifier deliberately does not import the exporter. It compares source XML,
resolved openpyxl values, provenance and exact Decimal sums, without recalculation.
"""
from __future__ import annotations

import argparse
from collections import Counter, defaultdict
from datetime import date, datetime, timezone
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
import hashlib
import json
from pathlib import Path
import posixpath
import re
import unicodedata
from xml.etree import ElementTree as ET
from zipfile import ZipFile

import openpyxl

NS = {"s": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
FIELDS = {
    "shipments": dict(zip("ABCDEFGHIJKLMNOPQRSTUVW", [
        "document_number", "month", "date", "customer_name", "manager_label",
        "payment_form", "product", "quantity_tonnes", "quantity_litres",
        "sale_price_per_tonne", "sale_price_per_litre", "customer_amount",
        "supplier_name", "purchase_price_unspecified_unit", "purchase_amount",
        "carrier_name", "transport_amount", "kvp_source", "profit_source",
        "paid_amount_source", "debt_overpayment_source", "term_source", "unlabelled_note"])),
    "payments": dict(zip("ABCDEFG", ["month", "date", "incoming_amount",
                                          "outgoing_amount", "purpose", "counterparty_name", "unlabelled_extra"])),
}
SOURCES = {"shipments": ("Бензовозы", 15, "ACDEFGHIJKLMNOPQRSTUVW"),
           "payments": ("Выписка", 2, "BCDEFG")}


def canonical_name(value):
    return " ".join(unicodedata.normalize("NFKC", str(value)).split()).casefold()


def sha256(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def independent_money(value, kind):
    if value is None or kind not in ("decimal", "string"):
        return None
    value = str(value).strip().replace("\u00a0", " ").replace("\u202f", " ")
    if kind == "string":
        decimal_lexeme = re.fullmatch(r"[+-]?(?:\d+(?:\.\d*)?|\.\d+)(?:[eE][+-]?\d+)?", value)
        if not decimal_lexeme:
            if not re.fullmatch(r"[+-]?(?:\d+|\d{1,3}(?: \d{3})+)(?:,\d{1,2})?", value):
                return None
            value = value.replace(" ", "").replace(",", ".")
    try:
        amount = Decimal(value)
        return amount if amount.is_finite() else None
    except InvalidOperation:
        return None


class Checks:
    def __init__(self):
        self.checked = Counter()
        self.failed = Counter()
        self.failures = []

    def check(self, condition, category, location):
        self.checked[category] += 1
        if not condition:
            self.failed[category] += 1
            if len(self.failures) < 100:
                self.failures.append({"category": category, "location": str(location)})

    def equal(self, actual, expected, category, location):
        self.check(actual == expected, category, location)


def source_snapshot(path):
    """Resolve shared strings/styles via openpyxl; preserve XML numeric lexemes."""
    formulas = openpyxl.load_workbook(path, data_only=False)
    cached = openpyxl.load_workbook(path, data_only=True)
    result = {}
    with ZipFile(path) as archive:
        relations = ET.fromstring(archive.read("xl/_rels/workbook.xml.rels"))
        paths = {item.attrib["Id"]: item.attrib["Target"] for item in relations}
        workbook = ET.fromstring(archive.read("xl/workbook.xml"))
        for sheet_node in workbook.find("s:sheets", NS):
            name = sheet_node.attrib["name"]
            target = paths[sheet_node.attrib["{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id"]]
            target = target.lstrip("/") if target.startswith("/") else posixpath.normpath("xl/" + target)
            root = ET.fromstring(archive.read(target))
            ws, cache = formulas[name], cached[name]
            nodes = {node.attrib["r"]: node for node in root.findall(".//s:sheetData/s:row/s:c", NS)}
            addresses = {ref for ref, node in nodes.items() if any(
                node.find("s:" + tag, NS) is not None for tag in ("v", "f", "is"))}
            # Include empty cells carrying comments or hyperlinks as source evidence.
            for row in ws:
                for cell in row:
                    if cell.comment or cell.hyperlink:
                        addresses.add(cell.coordinate)
            cells = {}
            for ref in sorted(addresses):
                node = nodes.get(ref)
                xml_type = node.attrib.get("t", "n") if node is not None else "n"
                v = node.find("s:v", NS) if node is not None else None
                f = node.find("s:f", NS) if node is not None else None
                cell = ws[ref]
                chosen = cache[ref] if f is not None else cell
                value = chosen.value
                kind = "empty"
                if isinstance(value, (datetime, date)):
                    kind, value = "datetime", value.isoformat()
                elif isinstance(value, bool):
                    kind = "boolean"
                elif chosen.data_type == "e":
                    kind = "error"
                elif xml_type == "n" and v is not None and v.text is not None:
                    kind, value = "decimal", v.text
                elif value is not None:
                    kind, value = "string", str(value)
                elif xml_type == "str" and v is not None:
                    kind, value = "string", ""
                cells[ref] = {
                    "value": value, "value_type": kind,
                    "formula": str(cell.value) if f is not None else None,
                    "has_formula": f is not None, "xml_type": xml_type,
                    "xml_value": v.text if v is not None else None,
                    "xml_value_present": v is not None,
                    "formula_xml": f.text if f is not None else None,
                    "formula_attributes": dict(f.attrib) if f is not None else None,
                    "number_format": cell.number_format,
                }
                if f is not None:
                    cells[ref]["cache_status"] = "error" if kind == "error" else "present" if value is not None else "empty_or_missing"
                if cell.comment:
                    cells[ref]["comment"] = {"text": cell.comment.text, "author": cell.comment.author}
                if cell.hyperlink:
                    cells[ref]["hyperlink"] = {"target": cell.hyperlink.target, "location": cell.hyperlink.location}
            result[name] = {"cells": cells, "state": ws.sheet_state,
                            "merged_ranges": sorted(str(r) for r in ws.merged_cells.ranges)}
    formulas.close()
    cached.close()
    return result


def business_rows(cells, start, columns):
    rows = defaultdict(dict)
    for address, value in cells.items():
        col, row = re.fullmatch(r"([A-Z]+)(\d+)", address).groups()
        if int(row) >= start:
            rows[int(row)][col] = value
    active, excluded = {}, []
    for row, row_cells in sorted(rows.items()):
        if any(col in columns and not cell["has_formula"] and cell["value"] is not None
               and str(cell["value"]).strip() for col, cell in row_cells.items()):
            active[row] = row_cells
        else:
            excluded.append(row)
    return active, excluded


def check_roles(path, checks, shipments):
    roles = json.loads(path.read_text(encoding="utf-8"))
    checks.equal(roles.get("status"), "specification_only", "roles", "status")
    checks.equal(roles.get("authorization_implemented"), False, "roles", "authorization_implemented")
    checks.equal(roles.get("default_decision"), "deny", "roles", "default_decision")
    for field in ("accounts", "role_assignments", "ownership_assignments", "additional_record_grants"):
        checks.equal(roles.get(field), [], "roles", field)
    definitions = roles.get("roles", [])
    by_id = {role["id"]: role for role in definitions}
    checks.equal(len(by_id), len(definitions), "roles", "unique ids")
    checks.check({"director", "admin", "manager"}.issubset(by_id), "roles", "requested roles")
    checks.check(any(p.get("resource") == "*" and "*" in p.get("actions", []) and p.get("scope") == "all"
                     for p in by_id.get("director", {}).get("permissions", [])), "roles", "director full access")
    scopes = roles.get("scope_definitions", {})
    for role in definitions:
        for permission in role.get("permissions", []) + role.get("proposed_permissions", []):
            checks.check(permission.get("scope") in scopes, "roles", f"{role['id']} scope")
            checks.check(bool(permission.get("resource")) and bool(permission.get("actions")), "roles", f"{role['id']} permission")
        if not role.get("enabled_for_assignment"):
            checks.equal(role.get("permissions"), [], "roles", f"{role['id']} disabled permissions")
    manager = by_id.get("manager", {})
    checks.check(all(p.get("scope", "").startswith("own_") for p in manager.get("permissions", [])), "roles", "manager ownership scopes")
    checks.check({"owner_user_id", "roles", "permissions"}.issubset(manager.get("forbidden_mutation_fields", [])), "roles", "manager cannot change privileges")
    ownership = roles.get("ownership_policy", {})
    checks.equal(ownership.get("imported_owner_default"), None, "roles", "imported_owner_default")
    checks.equal(ownership.get("infer_owners_from_source"), False, "roles", "no inferred owners")
    link = ownership.get("shipment_customer_field")
    checks.check(link in ("customer_id", "company_id"), "roles", "shipment company link")
    checks.check(all(link in item for item in shipments), "roles", "shipment company link exists in export")
    return {"path": str(path.resolve()), "sha256": sha256(path), "authorization_implemented": roles.get("authorization_implemented")}


def verify(source, export, roles_path):
    checks = Checks()
    source_hash = sha256(source)
    manifest = json.loads((export / "manifest.json").read_text(encoding="utf-8"))
    checks.equal(manifest["meta"]["source_sha256"], source_hash, "hashes", "manifest source")
    files = manifest.get("files", {})
    actual_files = {p.name for p in export.glob("*.json") if p.name != "manifest.json"}
    checks.equal(set(files), actual_files, "hashes", "complete file manifest")
    datasets = {}
    for filename, digest in files.items():
        checks.check(Path(filename).name == filename, "hashes", "safe manifest path")
        path = export / filename
        checks.equal(sha256(path), digest["sha256"], "hashes", filename)
        checks.equal(path.stat().st_size, digest["bytes"], "hashes", filename + " size")
        payload = json.loads(path.read_text(encoding="utf-8"))
        checks.check("meta" in payload and "data" in payload, "schema", filename)
        checks.equal(payload["meta"], manifest["meta"], "provenance", filename + " metadata")
        checks.equal(payload["meta"].get("google_verified"), False, "provenance", filename + " Google status")
        datasets[path.stem] = payload["data"]
    expected = source_snapshot(source)
    raw_sheets = datasets["raw_workbook"]["sheets"]
    raw = {sheet["name"]: sheet for sheet in raw_sheets}
    checks.equal(len(raw), len(raw_sheets), "source_cells", "unique sheets")
    checks.equal(set(raw), set(expected), "source_cells", "sheet names")
    for name, sheet in expected.items():
        actual = raw.get(name, {})
        checks.equal(actual.get("state"), sheet["state"], "sheet_metadata", name + " state")
        checks.equal(sorted(actual.get("merged_ranges", [])), sheet["merged_ranges"], "sheet_metadata", name + " merges")
        cells = actual.get("cells", {})
        checks.equal(set(cells), set(sheet["cells"]), "source_cells", name + " addresses including comments")
        for ref, cell in sheet["cells"].items():
            output_cell = cells.get(ref, {})
            for field, value in cell.items():
                checks.equal(output_cell.get(field), value, "source_cells", f"{name}!{ref} {field}")
            for field in ("comment", "hyperlink"):
                checks.equal(output_cell.get(field), cell.get(field), "source_cells", f"{name}!{ref} {field}")

    companies = datasets["companies"]
    by_company = {item["id"]: item for item in companies}
    checks.equal(len(by_company), len(companies), "links", "unique company ids")
    checks.equal(len({item["normalized_name"] for item in companies}), len(companies), "links", "unique normalized companies")
    by_name = {item["normalized_name"]: item for item in companies}
    expected_occurrences = defaultdict(list)
    all_records, rows_by_dataset = {}, {}

    def source_company(name, sheet, cell, role, record_id=None):
        if name is None or not str(name).strip():
            return None
        normalized = canonical_name(name)
        company = by_name.get(normalized)
        checks.check(company is not None, "company_provenance", f"{sheet}!{cell} missing company")
        if company is None:
            return None
        expected_occurrences[company["id"]].append({"sheet": sheet, "cell": cell, "role": role, "record_id": record_id})
        return company["id"]

    report = datasets["validation_report"]
    monetary = {}
    for dataset, (sheet, start, columns) in SOURCES.items():
        source_cells = expected[sheet]["cells"]
        rows, excluded = business_rows(source_cells, start, columns)
        rows_by_dataset[dataset] = rows
        items = datasets[dataset]
        actual_rows = [item["source"]["row"] for item in items]
        checks.equal(sorted(actual_rows), sorted(rows), "transaction_coverage", dataset)
        checks.equal(len(set(actual_rows)), len(items), "transaction_coverage", dataset + " unique rows")
        checks.equal(report["excluded_from_transaction_lists"][dataset[:-1] + "_template_rows"], excluded,
                     "transaction_coverage", dataset + " excluded templates")
        for item in items:
            row = item["source"]["row"]
            location = f"{sheet}!row{row}"
            checks.equal(item["source"], {"sheet": sheet, "row": row, "workbook_sha256": source_hash}, "provenance", location)
            checks.equal(item.get("owner_user_id", "missing"), None, "ownership", location)
            checks.check(item["id"] not in all_records, "links", location + " unique id")
            all_records[item["id"]] = item
            row_cells = rows.get(row, {})
            checks.equal(item["cells"], {col + str(row): raw[sheet]["cells"][col + str(row)] for col in row_cells}, "provenance", location + " complete cells")
            for col, field in FIELDS[dataset].items():
                cell = row_cells.get(col, {})
                value = cell.get("value")
                if dataset == "shipments" and field == "manager_label" and (cell.get("value_type") != "string" or not value or not str(value).strip()):
                    value = None
                checks.equal(item["fields"].get(field), value, "field_mapping", location + " " + field)
            links = [("D", "customer_id", "customer"), ("M", "supplier_id", "supplier"), ("P", "carrier_id", "carrier")] if dataset == "shipments" else [("F", "counterparty_id", "payment_counterparty")]
            for col, field, role in links:
                company_id = source_company(row_cells.get(col, {}).get("value"), sheet, col + str(row), role, item["id"])
                checks.equal(item.get(field), company_id, "links", location + " " + field)
            if dataset == "payments":
                for col, field in (("C", "incoming_amount"), ("D", "outgoing_amount")):
                    cell = row_cells.get(col, {})
                    amount = independent_money(cell.get("value"), cell.get("value_type"))
                    normalized = item.get("normalized_amounts", {}).get(field)
                    checks.equal(Decimal(normalized) if normalized is not None else None, amount, "money_normalization", location + " " + field)

        numeric_fields = ["quantity_tonnes", "quantity_litres", "customer_amount", "purchase_amount", "transport_amount", "paid_amount_source", "profit_source"] if dataset == "shipments" else ["incoming_amount", "outgoing_amount"]
        monetary[dataset] = {}
        for field in numeric_fields:
            col = next(col for col, name in FIELDS[dataset].items() if name == field)
            values = [Decimal(cells[col]["xml_value"]) for cells in rows.values() if cells.get(col, {}).get("value_type") == "decimal"]
            total = sum(values, Decimal(0))
            reported = report["source_totals"][dataset][field]
            checks.equal(Decimal(reported["total"]), total, "decimal_totals", dataset + " " + field)
            checks.equal(reported["numeric_cells"], len(values), "decimal_totals", dataset + " " + field + " count")
            checks.equal(reported["empty_or_non_numeric_cells"], len(rows) - len(values), "decimal_totals", dataset + " " + field + " nonnumeric count")
            monetary[dataset][field] = {"exact": str(total), "rounded_2dp": str(total.quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)), "numeric_cells": len(values)}
    for col, field in (("C", "incoming_amount"), ("D", "outgoing_amount")):
        total = sum((independent_money(cells.get(col, {}).get("value"), cells.get(col, {}).get("value_type")) or Decimal(0) for cells in rows_by_dataset["payments"].values()), Decimal(0))
        checks.equal(Decimal(report["normalized_payment_totals"][field]), total, "decimal_totals", "normalized " + field)
        monetary.setdefault("normalized_payments", {})[field] = {"exact": str(total), "rounded_2dp": str(total.quantize(Decimal("0.01"), rounding=ROUND_HALF_UP))}

    summary_refs = [("D5", "prepayment_summary", ["E5", "F5"])]
    summary_refs += [(f"{col}{row}", "supplier_summary", [f"{amount_col}{row}"]) for col, amount_col in (("M", "N"), ("P", "R"), ("T", "V")) for row in range(2, 12)]
    expected_summaries = []
    for ref, role, related in summary_refs:
        cells = expected["Бензовозы"]["cells"]
        cid = source_company(cells.get(ref, {}).get("value"), "Бензовозы", ref, role)
        if cid:
            expected_summaries.append({"counterparty_id": cid, "source": {"sheet": "Бензовозы", "cell": ref}, "role": role,
                                       "cells": {a: raw["Бензовозы"]["cells"].get(a) for a in [ref, *related]}})
    checks.equal(datasets["company_summaries"], expected_summaries, "summary_provenance", "company summaries")
    for row, cells in rows_by_dataset["payments"].items():
        extra = cells.get("G", {})
        if extra.get("value_type") == "string" and extra.get("value"):
            source_company(extra["value"], "Выписка", f"G{row}", "unlabelled_payment_note")

    stock_expected = []
    stock_fields = ["incoming_litres", "incoming_amount", "outgoing_litres", "outgoing_amount", "balance_litres", "balance_amount"]
    for title, month, cols, start in [("B1", "A", "BCDEFG", 4), ("J1", "I", "JKLMNO", 4), ("B16", "A", "BCDEFG", 19), ("J16", "I", "JKLMNO", 19), ("B31", "A", "BCDEFG", 34)]:
        cells = expected["Склад"]["cells"]
        label = cells.get(title, {}).get("value")
        cid = source_company(label, "Склад", title, "stock_summary")
        for row in range(start, start + 12):
            refs = [col + str(row) for col in cols]
            stock_expected.append({"counterparty_id": cid, "label": label, "month": cells.get(month + str(row), {}).get("value"),
                                   "source": {"sheet": "Склад", "row": row, "title_cell": title},
                                   "fields": dict(zip(stock_fields, [cells.get(ref, {}).get("value") for ref in refs])),
                                   "cells": {ref: raw["Склад"]["cells"].get(ref) for ref in refs}, "value_basis": "xlsx_cached_not_recalculated"})
    checks.equal(datasets["stock_summaries"], stock_expected, "summary_provenance", "stock summaries")

    def sorted_occurrences(items):
        return sorted(items, key=lambda x: (x["sheet"], x["cell"], x["role"], x.get("record_id") or ""))

    for company in companies:
        cid = company["id"]
        occurrences = expected_occurrences.get(cid, [])
        checks.check(bool(occurrences), "company_provenance", cid + " source exists")
        checks.equal(sorted_occurrences(company["occurrences"]), sorted_occurrences(occurrences), "company_provenance", cid + " occurrences")
        names = {str(expected[o["sheet"]]["cells"][o["cell"]]["value"]) for o in occurrences}
        checks.equal(set(company["source_names"]), names, "company_provenance", cid + " names")
        checks.equal(set(company["source_roles"]), {o["role"] for o in occurrences}, "company_provenance", cid + " roles")
        checks.equal(company.get("owner_user_id", "missing"), None, "ownership", cid)
        checks.equal(company.get("verification_status"), "source_only_not_registry_verified", "company_provenance", cid + " verification status")
        for field in ("inn", "kpp", "ogrn", "legal_address"):
            checks.equal(company.get(field), None, "company_provenance", cid + " unverified " + field)
        for dataset, role in (("payments", "payment_counterparty"), ("shipments", None)):
            ids = {o["record_id"] for o in occurrences if o.get("record_id") and ((o["role"] == role) if role else (o["role"] != "payment_counterparty"))}
            checks.equal(set(company[dataset[:-1] + "_ids"]), ids, "links", cid + " " + dataset)

    manager_items = datasets["manager_labels"]
    checks.equal(len({m["id"] for m in manager_items}), len(manager_items), "links", "unique manager ids")
    actual_managers = {}
    for manager in manager_items:
        checks.equal(manager.get("user_id", "missing"), None, "ownership", "manager " + manager["id"])
        checks.equal(manager.get("is_user_account"), False, "ownership", "manager label is not account")
        actual_managers[canonical_name(manager["source_label"])] = set(manager["shipment_ids"])
    expected_managers = defaultdict(set)
    for item in datasets["shipments"]:
        label = item["fields"].get("manager_label")
        if label:
            expected_managers[canonical_name(label)].add(item["id"])
    checks.equal(actual_managers, dict(expected_managers), "links", "manager shipment relationships")

    expected_counts = {"sheets": len(expected), "raw_cells": sum(len(s["cells"]) for s in expected.values()),
                       "formula_cells": sum(c["has_formula"] for s in expected.values() for c in s["cells"].values()),
                       "counterparties": len(companies), "shipment_rows": len(rows_by_dataset["shipments"]),
                       "payment_rows": len(rows_by_dataset["payments"]), "manager_labels": len(expected_managers),
                       "stock_monthly_rows": len(stock_expected), "company_summary_rows": len(expected_summaries)}
    payment_count = sum(any(independent_money(cells.get(col, {}).get("value"), cells.get(col, {}).get("value_type")) is not None for col in "CD") for cells in rows_by_dataset["payments"].values())
    expected_counts.update(payments_with_amount=payment_count, incomplete_payment_rows=len(rows_by_dataset["payments"]) - payment_count)
    for key, value in expected_counts.items():
        checks.equal(report["counts"].get(key), value, "counts", key)
        checks.equal(manifest["counts"].get(key), value, "counts", "manifest " + key)
    expected_issues = Counter()
    for sheet in expected.values():
        for cell in sheet["cells"].values():
            if cell["formula"] and "#REF!" in cell["formula"].upper():
                expected_issues["broken_formula_reference"] += 1
            if cell["value_type"] == "error":
                expected_issues["cached_excel_error"] += 1
            if cell["formula"] and cell["value"] is None:
                expected_issues["formula_cache_empty_or_missing"] += 1
    for key, value in expected_issues.items():
        checks.equal(report["issue_counts"].get(key), value, "counts", key)
    checks.equal(Counter(i["code"] for i in report["cell_issues"]), Counter(report["issue_counts"]), "counts", "issue details match counts")
    roles_result = check_roles(roles_path, checks, datasets["shipments"])
    return {"status": "passed" if not checks.failed else "failed", "checked_at_utc": datetime.now(timezone.utc).isoformat(),
            "source": {"file": str(source.resolve()), "sha256": source_hash}, "export": str(export.resolve()),
            "method": "Independent source ZIP/XML and openpyxl comparison; exact Decimal sums; no imported exporter code; no formula recalculation.",
            "google_verified": False, "registry_verified": False, "authorization_implemented": False,
            "checks_run": sum(checks.checked.values()), "checks_by_category": dict(checks.checked),
            "failed_checks": sum(checks.failed.values()), "failures_by_category": dict(checks.failed),
            "failures": checks.failures, "failure_details_truncated": sum(checks.failed.values()) > len(checks.failures),
            "source_counts": expected_counts, "independent_totals": monetary, "roles": roles_result}


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--source", type=Path, required=True)
    parser.add_argument("--export", type=Path, required=True)
    parser.add_argument("--roles", type=Path, default=Path(__file__).resolve().parents[1] / "config/roles.json")
    parser.add_argument("--output", type=Path, default=Path(__file__).resolve().parents[1] / "qa/verification.json")
    args = parser.parse_args()
    try:
        result = verify(args.source, args.export, args.roles)
    except Exception as error:
        result = {"status": "failed", "error_type": type(error).__name__, "error": str(error),
                  "source": str(args.source), "export": str(args.export)}
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(json.dumps(result, ensure_ascii=False, indent=2, allow_nan=False) + "\n", encoding="utf-8")
    print(json.dumps({k: result[k] for k in ("status", "checks_run", "failed_checks", "failures_by_category", "error_type", "error") if k in result}, ensure_ascii=False))
    return 0 if result["status"] == "passed" else 1


if __name__ == "__main__":
    raise SystemExit(main())
